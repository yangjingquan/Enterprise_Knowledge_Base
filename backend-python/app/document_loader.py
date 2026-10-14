"""
文档解析与切分。

把 PDF / Word / Excel / PPT / 网页 / TXT / Markdown 文件转成纯文本，再切成约 500 字的小块（chunk），
每块带上文档 ID、文件名、页码等元数据，方便回答时标注出处。
"""
import os
from functools import lru_cache
from pathlib import Path
import unicodedata
import pdfplumber
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

# 文本切分器：
#   chunk_size=500     每块最多约 500 个字符（太长检索不精准，太短上下文不完整）
#   chunk_overlap=50   相邻块重叠 50 字，避免一句话被切断后丢失语义
#   keep_separator="end"  分隔符（如句号）保留在上一块末尾，避免下一块以“。”开头
#   separators         按优先级尝试切分：段落 → 换行 → 中文句末标点 → 空格 → 任意字符
SPLITTER = RecursiveCharacterTextSplitter(
    chunk_size=500, chunk_overlap=50, keep_separator="end",
    separators=["\n\n", "\n", "。", "！", "？", "；", " ", ""],
)

# 部分 PDF 字体会把常用汉字输出成“CJK 部首补充”区的字符（如 部⻔、⻋辆），
# 外观一样但编码不同，会导致检索匹配不上。NFKC 处理不了这一区，这里手动映射回标准汉字。
_RADICAL_MAP = str.maketrans({
    "⻔": "门", "⻓": "长", "⻛": "风", "⻅": "见", "⺠": "民", "⻋": "车", "⻉": "贝",
    "⻜": "飞", "⻝": "食", "⻬": "齐", "⻢": "马", "⻚": "页", "⻰": "龙", "⻥": "鱼",
    "⻦": "鸟", "⻩": "黄", "⻨": "麦", "⻮": "齿", "⻆": "角", "⻘": "青", "⻙": "韦",
    "⻄": "西", "⻣": "骨", "⻤": "鬼", "⻢": "马", "⻧": "卤", "⻓": "长", "⻭": "齿",
})


def _table_to_md(table):
    """把表格（二维列表）转成 Markdown 表格文本，保留行列结构，让大模型能读懂表格。"""
    rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in table if r]
    if not rows:
        return ""
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([head, sep, *body])

# Excel 每组记录的最大字符数：略小于切分长度 500，保证一组正好落在一个块里
XLSX_BLOCK_CHARS = 450


def _cell_str(v):
    """把单元格的值转成字符串：None 变空串，整数形式的浮点数去掉“.0”。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------- 扫描件 OCR ----------

# 一页可提取的文字少于该字数，且页面上有图片，就判定为扫描页，转 OCR
OCR_MIN_CHARS = int(os.getenv("OCR_MIN_CHARS", "20"))
# 渲染分辨率：200 DPI 在识别率和速度之间比较均衡
OCR_DPI = int(os.getenv("OCR_DPI", "200"))
# 是否启用 OCR
OCR_ENABLED = os.getenv("OCR_ENABLED", "true").lower() in ("1", "true", "yes")


@lru_cache
def _ocr_engine():
    """
    RapidOCR（基于 onnxruntime 的 PaddleOCR 模型），模型文件随安装包自带（约 15MB），
    不需要联网下载，CPU 上运行，内存占用约 300~500MB。只在遇到扫描页时才加载。
    """
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR()


def _ocr_page(page) -> str:
    """
    把 PDF 页面渲染成图片后做文字识别，并按位置把识别结果还原成“行”：
    纵坐标相近（相差小于半个字高）的文字框归为同一行，行内按横坐标排序，用空格分隔。
    这样表格类扫描件（如点检记录表）也能保留“项目 标准 结果”的对应关系。
    """
    import numpy as np
    img = page.to_image(resolution=OCR_DPI).original.convert("RGB")
    result, _ = _ocr_engine()(np.array(img))
    if not result:
        return ""
    boxes = []
    for box, text, conf in result:
        if float(conf) < 0.5:   # 置信度太低的多为噪点
            continue
        ys = [pt[1] for pt in box]
        xs = [pt[0] for pt in box]
        boxes.append((sum(ys) / 4, min(xs), max(ys) - min(ys), text))
    boxes.sort()
    lines, current, last_y, last_h = [], [], None, 0
    for y, x, h, text in boxes:
        if last_y is not None and abs(y - last_y) > max(last_h, h) * 0.5:
            lines.append(current)
            current = []
        current.append((x, text))
        last_y, last_h = y, h
    if current:
        lines.append(current)
    return "\n".join(" ".join(t for _, t in sorted(line)) for line in lines)


def _load_xlsx(p: Path):
    """
    Excel：每个工作表作为一“页”（页码 = 工作表序号，从 1 开始），元数据带工作表名。
    每行数据转成“列名：值”的记录文本，例如
        备件编号：SP-SK-003；备件名称：导轨润滑油滤芯；规格型号：LF-10；含税单价(元)：85 ……
    为什么不直接用 Markdown 表格：表格块里大部分字符是重复的表头和“|”，
    向量检索时各块看起来几乎一样，按“LF-10”这类具体内容很难命中；
    记录文本每行都自带列名，语义更集中，大模型也能准确对应每个值的含义。
    多行按顺序累加成一组（每组不超过 XLSX_BLOCK_CHARS 个字符），每组前加工作表名。
    """
    from openpyxl import load_workbook
    # data_only=True：读取公式的计算结果而不是公式本身；read_only 节省内存
    wb = load_workbook(p, data_only=True, read_only=True)
    out = []
    for i, ws in enumerate(wb.worksheets, 1):
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = [_cell_str(v) for v in r]
            if any(cells):   # 跳过整行为空的行
                rows.append(cells)
        if not rows:
            continue
        # 第一行作为表头；表头为空的列用“第N列”代替
        header = [h or f"第{j + 1}列" for j, h in enumerate(rows[0])]
        records = []
        for row in rows[1:]:
            pairs = [f"{header[j]}：{v}" for j, v in enumerate(row) if v and j < len(header)]
            if pairs:
                records.append("；".join(pairs) + "。")
        if not records:
            records = ["表头：" + "、".join(header)]

        # 按长度分组，每组前加工作表名，让向量检索也能匹配到“哪张表”
        prefix = f"工作表：{ws.title}\n"
        blocks = []
        group = []
        for rec in records:
            if group and len(prefix) + len("\n".join(group + [rec])) > XLSX_BLOCK_CHARS:
                blocks.append(prefix + "\n".join(group))
                group = []
            group.append(rec)
        if group:
            blocks.append(prefix + "\n".join(group))
        out.append((i, "\n\n".join(blocks), {"sheet_name": ws.title}))
    wb.close()
    return out


def _shape_texts(shape):
    """递归提取一个形状里的文字：普通文本框、表格、组合形状里的子形状。"""
    texts = []
    if shape.shape_type == 6 and hasattr(shape, "shapes"):   # 6 = 组合形状
        for sub in shape.shapes:
            texts += _shape_texts(sub)
    elif getattr(shape, "has_table", False) and shape.has_table:
        rows = [[c.text for c in r.cells] for r in shape.table.rows]
        texts.append(_table_to_md(rows))
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            texts.append(t)
    return texts


def _load_pptx(p: Path):
    """PPT：每张幻灯片作为一页（页码 = 幻灯片序号），包含标题、正文、表格和演讲者备注。"""
    from pptx import Presentation
    prs = Presentation(p)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        parts = [f"幻灯片{i}：{title}" if title else f"幻灯片{i}"]
        for shape in slide.shapes:
            if slide.shapes.title is not None and shape.shape_id == slide.shapes.title.shape_id:
                continue   # 标题已经放在最前面，不重复
            parts += _shape_texts(shape)
        # 演讲者备注里经常有讲解细节，也一并入库
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"备注：{notes}")
        meta = {"slide_title": title} if title else {}
        out.append((i, "\n".join(parts), meta))
    return out


def _load_html(p: Path):
    """网页（URL 导入时保存的 HTML）：用 trafilatura 去掉导航、广告等，只保留正文和表格。"""
    import trafilatura
    html = p.read_text(encoding="utf-8", errors="ignore")
    text = trafilatura.extract(html, include_tables=True, include_comments=False, favor_recall=True) or ""
    return [(1, text, {})]


def load_pages(path: str):
    """
    解析文件，返回 [(页码, 文本, 额外元数据)]。
    “页”的含义随格式不同：PDF 是页，Excel 是工作表，PPT 是幻灯片，其他格式统一为 1。
    表格统一转成 Markdown，保留行列结构。
    """
    p = Path(path)
    ext = p.suffix.lower().lstrip(".")

    # PDF：逐页提取正文和表格，页码从 1 开始；没有文字层的扫描页走 OCR
    if ext == "pdf":
        out = []
        with pdfplumber.open(p) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if OCR_ENABLED and len(text.strip()) < OCR_MIN_CHARS and page.images:
                    out.append((i, _ocr_page(page), {"ocr": True}))
                    continue
                tables = [_table_to_md(t) for t in page.extract_tables()]
                out.append((i, "\n\n".join([text, *[t for t in tables if t]]), {}))
        return out
    # Word：段落 + 表格，Word 没有可靠的页码概念，统一记为第 1 页
    if ext == "docx":
        d = DocxDocument(p)
        parts = [para.text for para in d.paragraphs if para.text.strip()]
        parts += [_table_to_md([[c.text for c in r.cells] for r in t.rows]) for t in d.tables]
        return [(1, "\n".join(parts), {})]
    if ext == "xlsx":
        return _load_xlsx(p)
    if ext == "pptx":
        return _load_pptx(p)
    if ext in ("html", "htm"):
        return _load_html(p)
    # 纯文本 / Markdown：直接读取
    if ext in ("txt", "md"):
        return [(1, p.read_text(encoding="utf-8", errors="ignore"), {})]
    raise ValueError(f"不支持的文件类型: {ext}")


def split_document(path: str, document_id: int, file_name: str):
    """解析文件并切分成带元数据的 Document 列表（LangChain 的文档对象）。"""
    chunks = []
    for page_no, text, extra in load_pages(path):
        # NFKC 规范化：PDF 里有时会把“⼀”“⽇”等康熙部首字符当作汉字，统一转成标准汉字
        text = unicodedata.normalize("NFKC", text).translate(_RADICAL_MAP)
        for piece in SPLITTER.split_text(text):
            if piece.strip():   # 跳过空白块
                meta = {
                    "document_id": document_id,
                    "file_name": file_name,
                    "page": page_no,
                    "chunk_index": len(chunks),
                }
                meta.update(extra)   # 工作表名、幻灯片标题等格式相关的元数据
                chunks.append(Document(page_content=piece, metadata=meta))
    return chunks
