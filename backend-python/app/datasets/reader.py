"""
读取 csv / xlsx / xls，统一成"sheet + 行迭代器 + 合并区域"的形式，供后续识别表头、类型。
- 不一次性把整张表读进内存：xlsx 用 openpyxl 只读模式流式读取，csv 逐行读取；
- xlsx 只读模式拿不到合并单元格，单独解析 sheet XML 里的 <mergeCell ref="..."/>；
- xls 用 xlrd（formatting_info=True 才能读到合并单元格），老格式最多 65536 行，整表读入可接受。
"""
import codecs
import csv
import datetime as dt
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

MAX_XLSX_UNZIPPED = 200 * 1024 * 1024   # 防 zip 炸弹：xlsx 解压后总大小上限

_XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class DatasetError(Exception):
    """可以直接展示给用户的错误"""


@dataclass
class Sheet:
    name: str
    hidden: bool
    merges: list = field(default_factory=list)   # [(r1, c1, r2, c2)]，0 基、含两端
    _rows_fn: object = None

    def rows(self):
        """每次调用都从头开始迭代，产出 list（原始值：str/int/float/datetime/None）"""
        return self._rows_fn()


# ---------- 坐标工具 ----------

def col_letter(c: int) -> str:
    s = ""
    c += 1
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


_REF = re.compile(r"^\s*([A-Za-z]{1,3})(\d+)\s*(?::\s*([A-Za-z]{1,3})(\d+))?\s*$")


def parse_range(ref: str):
    """'A5:H120' → (4, 0, 119, 7)，0 基含两端"""
    m = _REF.match(ref or "")
    if not m:
        raise DatasetError(f"范围格式不正确：{ref}，示例 A5:H120")
    c1, r1 = col_index(m.group(1)), int(m.group(2)) - 1
    if m.group(3):
        c2, r2 = col_index(m.group(3)), int(m.group(4)) - 1
    else:
        c2, r2 = c1, r1
    return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


def format_range(r1, c1, r2, c2) -> str:
    return f"{col_letter(c1)}{r1 + 1}:{col_letter(c2)}{r2 + 1}"


# ---------- 各格式读取 ----------

def check_file(path: str, ext: str):
    p = Path(path)
    head = p.open("rb").read(8)
    if ext == "xlsx":
        if head.startswith(_XLS_MAGIC):
            raise DatasetError("文件已加密或实际是 xls 格式，请取消密码或另存为 xlsx 后再上传")
        if not head.startswith(b"PK"):
            raise DatasetError("不是有效的 xlsx 文件")
        with zipfile.ZipFile(p) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_XLSX_UNZIPPED:
                raise DatasetError("文件解压后过大（超过 200MB），请拆分后上传")
    elif ext == "xls":
        if not head.startswith(_XLS_MAGIC):
            raise DatasetError("不是有效的 xls 文件")
    elif ext == "csv":
        if b"\x00" in p.open("rb").read(4096):
            raise DatasetError("不是有效的 csv 文本文件")
    else:
        raise DatasetError(f"不支持的数据文件类型：{ext}")


def _xlsx_merges(path) -> dict:
    """{sheet 名: [(r1,c1,r2,c2)]}：直接解析 workbook.xml 与 sheet XML，只读模式下 openpyxl 不提供合并信息"""
    out = {}
    with zipfile.ZipFile(path) as z:
        wb = z.read("xl/workbook.xml").decode("utf8", "ignore")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf8", "ignore")
        rid2target = dict(re.findall(r'<Relationship[^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', rels))
        rid2target.update({k: v for v, k in re.findall(r'<Relationship[^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
        for tag in re.findall(r"<sheet\b[^>]*/?>", wb):
            name = re.search(r'name="([^"]*)"', tag)
            rid = re.search(r'r:id="([^"]*)"', tag)
            if not name or not rid:
                continue
            target = rid2target.get(rid.group(1), "")
            target = target.lstrip("/")
            target = target if target.startswith("xl/") else "xl/" + target
            try:
                xml = z.read(target).decode("utf8", "ignore")
            except KeyError:
                continue
            from html import unescape
            out[unescape(name.group(1))] = [parse_range(r) for r in re.findall(r'<mergeCell[^>]*ref="([^"]+)"', xml)]
    return out


def _open_xlsx(path):
    import openpyxl
    merges = _xlsx_merges(path)
    sheets = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        def rows(ws_name=ws.title):
            w = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                for r in w[ws_name].iter_rows(values_only=True):
                    yield list(r)
            finally:
                w.close()
        sheets.append(Sheet(ws.title, ws.sheet_state != "visible", merges.get(ws.title, []), rows))
    wb.close()
    return sheets


def _open_xls(path):
    import xlrd
    book = xlrd.open_workbook(path, formatting_info=True, on_demand=False)
    sheets = []
    for sh in book.sheets():
        merges = [(r1, c1, r2 - 1, c2 - 1) for r1, r2, c1, c2 in sh.merged_cells]

        def rows(sh=sh):
            for r in range(sh.nrows):
                out = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        out.append(None)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            out.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                        except Exception:
                            out.append(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        v = cell.value
                        out.append(int(v) if float(v).is_integer() else v)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        out.append(bool(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        out.append(None)
                    else:
                        out.append(cell.value)
                yield out
        sheets.append(Sheet(sh.name, getattr(sh, "visibility", 0) != 0, merges, rows))
    return sheets


def detect_encoding(path) -> str:
    raw = Path(path).open("rb").read(256 * 1024)
    if raw.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    for enc in ("utf-8", "gb18030"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError as e:
            # 截断在多字节字符中间不算错
            if enc == "utf-8" and e.start >= len(raw) - 3:
                return enc
    return "gb18030"


def _open_csv(path):
    enc = detect_encoding(path)
    sample = Path(path).open(encoding=enc, errors="replace", newline="").read(64 * 1024)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","

    def rows():
        with open(path, encoding=enc, errors="replace", newline="") as f:
            for r in csv.reader(f, delimiter=delim):
                yield [v if v.strip() != "" else None for v in r]
    return [Sheet(Path(path).stem, False, [], rows)]


def open_workbook(path: str, ext: str):
    ext = ext.lower()
    check_file(path, ext)
    try:
        if ext == "xlsx":
            return _open_xlsx(path)
        if ext == "xls":
            return _open_xls(path)
        return _open_csv(path)
    except DatasetError:
        raise
    except Exception as e:
        raise DatasetError(f"文件解析失败：{e}") from e
