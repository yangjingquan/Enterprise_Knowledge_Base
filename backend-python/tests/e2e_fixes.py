"""
第 11 项修复的端到端测试（需先启动 Python 8000 与 Java 8080）。
python -m tests.e2e_fixes [http://localhost:8080] [Python 日志路径]
覆盖：网页导入 SSRF、按请求切换检索方式、同库去重、失败重试与原因、Word 章节、删除顺序（部分在外部脚本里做停服测试）。
"""
import io
import json
import sys
import time
from pathlib import Path

import requests
from docx import Document as Docx

from tests.kb_auth import login

BASE = (sys.argv[1] if len(sys.argv) > 1 else "http://localhost:8080") + "/api"
PYLOG = sys.argv[2] if len(sys.argv) > 2 else "/tmp/kbt/logs/py.log"
DEMO = Path(__file__).resolve().parents[2] / "demo-data"
results = []


def check(name, cond, info=""):
    results.append((name, bool(cond)))
    print(("PASS " if cond else "FAIL ") + name + (f"  [{info}]" if info and not cond else ""), flush=True)


def wait_doc(h, did, timeout=300):
    t0 = time.time()
    while time.time() - t0 < timeout:
        d = next((x for x in requests.get(f"{BASE}/documents", headers=h).json() if x["id"] == did), None)
        if d and d["status"] in ("ready", "failed"):
            return d
        time.sleep(1)
    return d


def make_docx():
    d = Docx()
    d.add_paragraph("本制度适用于全体员工。")
    d.add_heading("第一章 总则", level=1)
    d.add_paragraph("公司实行标准工时制度。")
    d.add_heading("1.1 工作时间", level=2)
    d.add_paragraph("每日工作时间为上午九点至下午六点，午休一小时，迟到超过三十分钟按旷工半天处理。")
    d.add_heading("第二章 休假", level=1)
    d.add_heading("2.1 年假", level=2)
    d.add_paragraph("员工入职满一年后享有带薪年假五天，紫藤花园区员工另加一天。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "工龄", "年假", "10年", "10天"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def ask(h, q, **kw):
    r = requests.post(f"{BASE}/qa/ask/stream", headers=h, json={"question": q, "history": [], **kw}, stream=True)
    r.encoding = "utf-8"
    ev = [json.loads(l[6:]) for l in r.iter_lines(decode_unicode=True) if l and l.startswith("data: ")]
    src = next((e["data"] for e in ev if e["type"] == "sources"), [])
    return ev, src


def main():
    h = login(BASE, "admin", "admin123")
    kbs = {k["name"]: k["id"] for k in requests.get(f"{BASE}/kbs", headers=h).json()}
    kb_a, kb_b = kbs["公共知识库"], kbs["人事制度"]
    created = []

    # ---------- 10. ping ----------
    r = requests.get(f"{BASE}/auth/ping")
    check("10 GET /api/auth/ping 无需登录 200", r.status_code == 200 and r.json() == {"status": "ok"}, r.text)

    # ---------- 2. SSRF ----------
    for u in ["http://127.0.0.1:8000/health", "http://192.168.1.1/", "http://169.254.169.254/latest/meta-data/",
              "http://[::1]:8080/", "http://localhost:8080/api/auth/ping", "http://10.0.0.1/"]:
        r = requests.post(f"{BASE}/documents/url", headers=h, json={"url": u, "kbId": kb_a})
        check(f"2 网页导入拒绝内网地址 {u}", r.status_code == 400 and "不允许" in r.json().get("error", ""), r.text[:200])
    # 公网 302 → 内网：httpbin 的 redirect-to
    r = requests.post(f"{BASE}/documents/url", headers=h,
                      json={"url": "https://httpbin.org/redirect-to?url=http%3A%2F%2F127.0.0.1%3A8000%2Fhealth", "kbId": kb_a})
    check("2 公网重定向到内网被拒绝", r.status_code == 400 and ("不允许" in r.text or "无法访问" in r.text), r.text[:200])
    r = requests.post(f"{BASE}/documents/url", headers=h, json={"url": "https://www.python.org/about/", "kbId": kb_a})
    ok = r.status_code == 200
    check("2 公网网页正常导入", ok, r.text[:200])
    if ok:
        created.append(r.json()["id"])
        d = wait_doc(h, r.json()["id"])
        check("2 公网网页入库成功", d["status"] == "ready", d)

    # ---------- 5. 去重 ----------
    pdf = (DEMO / "员工手册.pdf").read_bytes()
    r1 = requests.post(f"{BASE}/documents/upload", headers=h, files={"file": ("员工手册.pdf", pdf)}, data={"kbId": kb_a})
    check("5 首次上传成功", r1.status_code == 200, r1.text[:200])
    d1 = r1.json(); created.append(d1["id"])
    check("5 记录了 SHA-256", len(d1.get("fileHash") or "") == 64, d1)
    r2 = requests.post(f"{BASE}/documents/upload", headers=h, files={"file": ("员工手册-副本.pdf", pdf)}, data={"kbId": kb_a})
    check("5 同库重复上传被拒（409 文件已存在）", r2.status_code == 409 and "文件已存在" in r2.json().get("error", ""), r2.text)
    r3 = requests.post(f"{BASE}/documents/upload", headers=h, files={"file": ("员工手册.pdf", pdf)}, data={"kbId": kb_b})
    check("5 不同知识库允许上传相同文件", r3.status_code == 200, r3.text[:200])
    if r3.status_code == 200:
        created.append(r3.json()["id"])
    d = wait_doc(h, d1["id"])
    check("5 首次上传的文档正常入库", d["status"] == "ready" and d["chunkCount"] > 0, d)

    # ---------- 4. 检索方式 ----------
    for m in ["vector", "hybrid", "rerank"]:
        mark = Path(PYLOG).stat().st_size
        ev, src = ask(h, "工作满10年有几天年假？", kbId=kb_a, mode="doc", retrievalMode=m)
        log = Path(PYLOG).read_text(errors="ignore")[mark:]
        check(f"4 流式问答 retrievalMode={m} 生效（日志）", f"requested_mode={m} mode={m}" in log, log[-300:])
        if m == "rerank":
            check("4 rerank 模式来源带 rerank_score", src and src[0]["detail"].get("rerank_score") is not None, src[:1])
        if m == "vector":
            check("4 vector 模式来源不含 BM25 名次", src and all(s["detail"].get("bm25_rank") is None for s in src), src[:1])
        check(f"4 {m} 模式正常结束", ev and ev[-1]["type"] == "done" and any(e["type"] == "saved" for e in ev))
    mark = Path(PYLOG).stat().st_size
    ask(h, "工作满10年有几天年假？", kbId=kb_a, mode="doc")
    log = Path(PYLOG).read_text(errors="ignore")[mark:]
    check("4 不指定时使用默认方式", "requested_mode=- mode=rerank" in log, log[-300:])
    mark = Path(PYLOG).stat().st_size
    ask(h, "工作满10年有几天年假？", kbId=kb_a, mode="doc", retrievalMode="evil'; drop")
    log = Path(PYLOG).read_text(errors="ignore")[mark:]
    check("4 非法取值按默认处理", "requested_mode=- mode=rerank" in log, log[-300:])
    r = requests.post(f"{BASE}/qa/ask", headers=h, json={"question": "婚假有多少天？", "kbId": kb_a, "mode": "doc",
                                                          "retrievalMode": "hybrid"})
    check("4 非流式问答也支持 retrievalMode", r.status_code == 200 and r.json().get("sources"), r.text[:200])

    # ---------- 6. 失败重试 ----------
    bad = b"%PDF-1.4\n" + b"\x00garbage" * 200
    r = requests.post(f"{BASE}/documents/upload", headers=h, files={"file": ("损坏文件.pdf", bad)}, data={"kbId": kb_a})
    did = r.json()["id"]; created.append(did)
    d = wait_doc(h, did)
    check("6 损坏文件处理失败", d["status"] == "failed", d)
    check("6 记录失败原因", bool(d.get("errorMsg")), d)
    print("   失败原因：", d.get("errorMsg"))
    r = requests.post(f"{BASE}/documents/{did}/retry", headers=h)
    check("6 失败文档可重试", r.status_code == 200 and r.json()["status"] in ("pending", "processing", "failed"), r.text[:200])
    d = wait_doc(h, did)
    check("6 重试后仍失败（原文件仍损坏），原因保留", d["status"] == "failed" and d.get("errorMsg"), d)
    # 把服务器上的原文件替换成正常 PDF，模拟"修复了服务/依赖后重试"
    Path(d["filePath"]).write_bytes((DEMO / "差旅与费用报销管理制度.pdf").read_bytes())
    requests.post(f"{BASE}/documents/{did}/retry", headers=h)
    d = wait_doc(h, did)
    check("6 重试成功后变为 ready 且清除原因", d["status"] == "ready" and not d.get("errorMsg") and d["chunkCount"] > 0, d)
    r = requests.post(f"{BASE}/documents/{did}/retry", headers=h)
    check("6 非失败状态不能重试（409）", r.status_code == 409, r.text)
    hr = login(BASE, "hr01", "Hr@2026")
    r = requests.post(f"{BASE}/documents/{did}/retry", headers=hr)
    check("6 无权限用户重试 → 403/404", r.status_code in (403, 404), r.status_code)

    # ---------- 7. Word 章节 ----------
    r = requests.post(f"{BASE}/documents/upload", headers=h, files={"file": ("考勤与休假制度.docx", make_docx())}, data={"kbId": kb_a})
    wid = r.json()["id"]; created.append(wid)
    d = wait_doc(h, wid)
    check("7 Word 入库成功", d["status"] == "ready", d)
    ev, src = ask(h, "紫藤花园区员工年假有几天？", documentId=wid)
    s0 = src[0] if src else {}
    check("7 Word 引用带章节路径", s0.get("section") == "第二章 休假 > 2.1 年假", s0)
    ctx = requests.get(f"{BASE}/documents/{wid}/context", headers=h, params={"chunkIndex": s0.get("chunk_index", 0)}).json()
    check("7 上下文接口返回章节", any(c.get("section") for c in ctx), ctx)

    # ---------- 9. 删除顺序（正常路径） ----------
    del_id = created.pop()   # Word 文档
    r = requests.delete(f"{BASE}/documents/{del_id}", headers=h)
    check("9 正常删除：cleanupPending=false", r.status_code == 200 and r.json().get("cleanupPending") is False, r.text)
    ev, src = ask(h, "紫藤花园区员工年假有几天？", kbId=kb_a, mode="doc")
    check("9 删除后检索不到该文档", all(s["document_id"] != del_id for s in src), src[:2])
    r = requests.post(f"{BASE}/documents/admin/purge-orphans", headers=hr, params={"dryRun": "true"})
    check("9 普通用户不能调用清理孤立向量（403）", r.status_code == 403, r.status_code)
    r = requests.post(f"{BASE}/documents/admin/purge-orphans", headers=h, params={"dryRun": "true"})
    check("9 管理员 dryRun 清理孤立向量", r.status_code == 200 and "orphanDocumentIds" in r.json(), r.text)

    Path("/tmp/kbt/e2e_fixes_created.json").write_text(json.dumps(created))
    failed = [n for n, ok in results if not ok]
    print(f"\n{len(results) - len(failed)}/{len(results)} passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
