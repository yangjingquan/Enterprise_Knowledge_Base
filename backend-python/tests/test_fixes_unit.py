"""Word 章节、检索方式参数、入库错误原因的单元测试"""
import io

import pytest
from docx import Document as Docx

from app import document_loader, main, rag, retrieval


def _docx(tmp_path):
    d = Docx()
    d.add_paragraph("前言内容，位于第一个标题之前。")
    d.add_heading("第一章 总则", level=1)
    d.add_paragraph("总则正文。")
    d.add_heading("1.1 适用范围", level=2)
    d.add_paragraph("适用于全体员工。")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "项目", "标准"
    d.add_heading("第二章 休假", level=1)
    d.add_paragraph("年假五天。")
    p = tmp_path / "a.docx"
    d.save(p)
    return p


def test_docx_sections(tmp_path):
    pages = document_loader.load_pages(str(_docx(tmp_path)))
    secs = [e.get("section") for _, _, e in pages]
    assert secs == [None, "第一章 总则", "第一章 总则 > 1.1 适用范围", "第二章 休假"]
    assert all(p == 1 for p, _, _ in pages)
    assert "| 项目 | 标准 |" in pages[2][1]            # 表格归入所在章节，且保持正文顺序
    assert pages[3][1].startswith("第二章 休假")       # 章节文本以标题开头


def test_docx_chunks_have_section(tmp_path):
    chunks = document_loader.split_document(str(_docx(tmp_path)), 1, "a.docx")
    assert chunks[0].metadata.get("section") is None
    assert {c.metadata.get("section") for c in chunks[1:]} == {"第一章 总则", "第一章 总则 > 1.1 适用范围", "第二章 休假"}
    assert [c.metadata["chunk_index"] for c in chunks] == list(range(len(chunks)))


def test_location_uses_section():
    assert rag._location({"file_name": "a.docx", "page": 1, "section": "第一章 > 1.1"}) == "章节「第一章 > 1.1」"
    assert rag._location({"file_name": "a.pdf", "page": 3}) == "第3页"


@pytest.mark.parametrize("v,exp", [("vector", "vector"), ("HYBRID", "hybrid"), ("rerank", "rerank"),
                                   (None, None), ("", None), ("bm25", None), ("vector; drop", None)])
def test_retrieval_mode_param(v, exp):
    req = main.AskReq(question="q", retrieval_mode=v)
    assert main._rmode(req) == exp


def test_retrieve_invalid_mode_uses_default(monkeypatch, capsys):
    monkeypatch.setattr(retrieval.vectorstore, "search", lambda *a, **k: [])
    monkeypatch.setattr(retrieval.bm25.index, "search", lambda *a, **k: [])
    retrieval.retrieve("q", 3, None, [1], "nonsense")
    assert f"mode={retrieval.DEFAULT_MODE if retrieval.DEFAULT_MODE != 'rerank' or retrieval.reranker.enabled() else 'hybrid'}" in capsys.readouterr().out


def test_ingest_error_detail(tmp_path):
    from fastapi.testclient import TestClient
    bad = tmp_path / "x.pdf"
    bad.write_bytes(b"not a pdf")
    c = TestClient(main.app)
    r = c.post("/ingest", json={"document_id": 1, "file_path": str(bad), "file_name": "x.pdf", "kb_id": 1})
    assert r.status_code == 422 and "解析失败" in r.json()["detail"]
    r = c.post("/ingest", json={"document_id": 1, "file_path": str(tmp_path / "none.pdf"), "file_name": "n.pdf", "kb_id": 1})
    assert r.status_code == 422 and "不存在" in r.json()["detail"]
    empty = tmp_path / "e.txt"
    empty.write_text("   ")
    r = c.post("/ingest", json={"document_id": 1, "file_path": str(empty), "file_name": "e.txt", "kb_id": 1})
    assert r.status_code == 422 and "没有从文件中提取到文字" in r.json()["detail"]
